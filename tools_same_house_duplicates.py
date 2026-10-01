# -*- coding: utf-8 -*-
"""One person entered twice in the same house, in the same round.

The strongest duplicate signal the record has. Neither the shape-matcher nor the
surname bucket finds these, because the two spellings differ - Catherine Lucy
Adams beside Catherine S Adam, George Milton Sydney Johnson beside George S M
Johnson - and the ages can differ too, because two transcribers read the same
figure differently.

**The test is deliberately strict**, because a loose one finds families rather
than duplicates: a first forename that matches outright, a surname within two
letters, and **ages within three years**. A father and his son share a house, a
round and a name, and they are decades apart; two readings of one line are not.
Gerald and Gertrude Kirk, both 63, are a husband and a wife and are not caught,
because the first forename has to match.

The pairs are printed for reading before anything is done, and the older record
is kept.

  railway run python3 tools_same_house_duplicates.py          # show
  railway run python3 tools_same_house_duplicates.py --apply  # fold them
"""
import os, sys, psycopg2

APPLY = '--apply' in sys.argv
EXACT = '--exact' in sys.argv       # only fold pairs whose names are letter for letter
SHEET = '--sheet' in sys.argv       # write the rest out for A. Hagues to confirm

FIND = """
SELECT a.person_id, pa.first_name, pa.last_name, a.age_at_census, a.id,
       b.person_id, pb.first_name, pb.last_name, b.age_at_census, b.id,
       a.census_year, a.property_id
  FROM census_entries a
  JOIN census_entries b ON b.census_year = a.census_year
                       AND b.property_id = a.property_id
                       AND b.person_id > a.person_id
  JOIN people pa ON pa.id = a.person_id
  JOIN people pb ON pb.id = b.person_id
 WHERE a.property_id IS NOT NULL
   AND LOWER(LEFT(pa.last_name, 4)) = LOWER(LEFT(pb.last_name, 4))
   AND ABS(LENGTH(pa.last_name) - LENGTH(pb.last_name)) <= 2
   -- the FIRST forename must match outright: "George Milton Sydney" and
   -- "George S M" are one man, but "Gerald" and "Gertrude" are a husband and
   -- a wife, and "George" and "Georgina" are a father and his daughter
   AND LOWER(SPLIT_PART(TRIM(pa.first_name), ' ', 1)) = LOWER(SPLIT_PART(TRIM(pb.first_name), ' ', 1))
   -- and they must be the same age, give or take a transcriber. A parent and a
   -- child of one name are decades apart; two readings of one line are not
   AND a.age_at_census IS NOT NULL AND b.age_at_census IS NOT NULL
   AND ABS(a.age_at_census - b.age_at_census) <= 3
   AND (pa.bio IS NULL OR pb.bio IS NULL)
 ORDER BY a.census_year, a.property_id, pa.last_name
"""


def main():
    c = psycopg2.connect(os.environ.get('DATABASE_PUBLIC_URL') or os.environ['DATABASE_URL'])
    cur = c.cursor()
    cur.execute(FIND)
    pairs = cur.fetchall()
    done, gone = set(), set()
    kept = 0
    doubtful = []
    for (aid, afn, aln, aage, arow, bid, bfn, bln, bage, brow, yr, prop) in pairs:
        if bid in done or bid in gone or aid in gone:
            continue   # a chained pair: one of these was folded away earlier in this run
        same = (afn.strip().lower(), aln.strip().lower()) == (bfn.strip().lower(), bln.strip().lower())
        if EXACT and not same:
            doubtful.append((yr, prop, aid, afn, aln, aage, bid, bfn, bln, bage))
            continue
        done.add(bid)
        print(f"  {yr} house {prop}: #{aid} {afn} {aln} ({aage})  <-  #{bid} {bfn} {bln} ({bage})")
        kept += 1
        if not APPLY:
            continue
        cur.execute("""INSERT INTO person_alias (person_id, first_name, last_name, born_year, made_by)
                       SELECT %s,%s,%s,(SELECT born_year FROM people WHERE id=%s),
                              'the same person under another spelling'
                       FROM (SELECT 1) t WHERE NOT EXISTS
                       (SELECT 1 FROM person_alias x
                         WHERE LOWER(TRIM(x.first_name))=LOWER(TRIM(%s))
                           AND LOWER(TRIM(x.last_name))=LOWER(TRIM(%s)))""",
                    (aid, bfn, bln, bid, bfn, bln))
        cur.execute("""UPDATE census_entries n SET person_id=%s
                        WHERE n.person_id=%s AND NOT EXISTS
                          (SELECT 1 FROM census_entries o
                            WHERE o.person_id=%s AND o.census_year=n.census_year
                              AND o.property_id IS NOT DISTINCT FROM n.property_id)""",
                    (aid, bid, aid))
        cur.execute("DELETE FROM census_entries WHERE person_id=%s", (bid,))
        cur.execute("DELETE FROM property_residents WHERE person_id=%s", (bid,))
        cur.execute("DELETE FROM people WHERE id=%s", (bid,))
        gone.add(bid)
    if gone:
        print(f"  (run again: chained pairs are skipped once, not resolved)")
    print(f"\n  {kept} pairs" + (", folded and committed" if APPLY else ". Add --apply."))
    if APPLY:
        c.commit()
    if SHEET and doubtful:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "same house, same age"
        ws.append(['year', 'house', 'keep (id)', 'name', 'age',
                   'fold away (id)', 'name', 'age', 'ONE PERSON? y/n'])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        fill = PatternFill('solid', fgColor='FFF2CC')
        ws['I1'].fill = fill
        for row in doubtful:
            ws.append(list(row[:5]) + list(row[5:]) + [''])
            ws[ws.max_row][8].fill = fill
        for i, w in enumerate((7, 8, 11, 30, 6, 13, 30, 6, 16), 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'
        out = os.path.expanduser('~/Desktop/Same house duplicates.xlsx')
        wb.save(out)
        print(f"  {len(doubtful)} pairs whose spellings differ written to {out}")


if __name__ == '__main__':
    main()
