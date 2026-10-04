"""The Peveril Drive run, laid out schedule by schedule so it can be matched to
the street from the page. Nine households have no house; eight houses have no
household; and two more households name addresses the property list does not
hold at all."""
import os, json, psycopg2
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

P = json.load(open('data/all_props.json'))
props = {p['id']: (p.get('address') or p.get('name')) for p in P}
c = psycopg2.connect(os.environ.get('DATABASE_PUBLIC_URL') or os.environ['DATABASE_URL'])
cur = c.cursor()

cur.execute("""SELECT ce.census_household_num, MIN(ce.property_id),
                      MIN(COALESCE(NULLIF(ce.address,''),ce.unresolved_address)),
                      STRING_AGG(p.first_name||' '||p.last_name||' ('||COALESCE(ce.age_at_census::text,'?')||')',
                                 '; ' ORDER BY ce.id),
                      STRING_AGG(DISTINCT COALESCE(ce.occupation_at_census,''), ' / ')
                 FROM census_entries ce JOIN people p ON p.id=ce.person_id
                WHERE ce.census_year=1939 AND ce.source ILIKE '%RMGB%'
                  AND ce.census_household_num BETWEEN 114 AND 130
                GROUP BY 1 ORDER BY 1""")
run = {r[0]: r for r in cur.fetchall()}

cur.execute("SELECT DISTINCT property_id FROM census_entries WHERE census_year=1939 AND property_id IS NOT NULL")
has = {r[0] for r in cur.fetchall()}
cur.execute("SELECT property_id FROM census_unoccupied WHERE census_year=1939")
emp = {r[0] for r in cur.fetchall()}
free = [p for p in P if 'peveril' in str(p.get('street') or '').lower()
        and p['id'] not in has and p['id'] not in emp]

wb = Workbook(); ws = wb.active; ws.title = "Peveril Drive 114-130"
cols = ['Schedule', 'House the record gives', 'Who is in it', 'What they do']
ws.append(cols)
for col in ws[1]:
    col.font = Font(bold=True, color="FFFFFF")
    col.fill = PatternFill("solid", fgColor="1F4E46")
for s in range(114, 131):
    r = run.get(s)
    if not r:
        ws.append([s, "— no household in the record —", "", ""]); continue
    house = props.get(r[1]) or r[2] or ''
    if not r[1]:
        house = "*** NO HOUSE *** " + (r[2] or '')
    ws.append([s, house, r[3], (r[4] or '').strip(' /')])
ws.append([]); ws.append(["Houses on Peveril Drive with nothing at all for 1939"])
ws.cell(ws.max_row, 1).font = Font(bold=True)
for p in sorted(free, key=lambda x: x['id']):
    ws.append(["", p.get('address') or p.get('name'), "", ""])
ws.append([]); ws.append(["Households naming a Peveril address the property list does not hold"])
ws.cell(ws.max_row, 1).font = Font(bold=True)
cur.execute("""SELECT ce.census_household_num, MIN(COALESCE(NULLIF(ce.address,''),ce.unresolved_address)),
                      STRING_AGG(p.first_name||' '||p.last_name,'; ' ORDER BY ce.id)
                 FROM census_entries ce JOIN people p ON p.id=ce.person_id
                WHERE ce.census_year=1939 AND ce.property_id IS NULL
                  AND COALESCE(NULLIF(ce.address,''),ce.unresolved_address) ILIKE '%peveril%'
                  AND ce.census_household_num NOT BETWEEN 114 AND 130
                GROUP BY 1 ORDER BY 1""")
for r in cur.fetchall(): ws.append([r[0], r[1], r[2], ""])

ws.freeze_panes = "A2"
for i, w in enumerate([10, 46, 64, 46], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.alignment = Alignment(vertical='top', wrap_text=True)
out = os.path.expanduser('~/Desktop/Peveril Drive - the run.xlsx')
wb.save(out)
print(f"{len([s for s in range(114,131) if run.get(s) and not run[s][1]])} households with no house")
print(f"{len(free)} houses with nothing")
print(out)
