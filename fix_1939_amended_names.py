# -*- coding: utf-8 -*-
"""The 1939 Register's amended surnames, put back the way the page had them.

The Register was not a census. It was an identity register, kept up to date for
decades, so a woman who married after 1939 has her married name written across
the original and the online indexes show the LATER name. Ivy Burton, cook at
Penrhyn House on 29 September 1939, is indexed as Ivy Luke because somebody
wrote LUKE over Burton years afterwards.

The rule this record follows: **a person is named as the page named them on the
night**. The amendment is a real and valuable fact - it dates a marriage - but
it is a fact about what happened later, not the name of the woman in the house.

So the name goes back to the one she had, the married name is kept as an alias
so a search for it still finds her and a re-import cannot make a second person,
and the census row says what the page does.

Usage:
  railway run python3 fix_1939_amended_names.py          # show
  railway run python3 fix_1939_amended_names.py --apply  # do it

Or take the names from the filled-in sheet (amended_names_worksheet.py writes it):
  railway run python3 fix_1939_amended_names.py --sheet "~/Desktop/1939 names - fill in.xlsx"
"""
import os, sys, psycopg2

APPLY = '--apply' in sys.argv
SHEET = next((a.split('=', 1)[1] for a in sys.argv if a.startswith('--sheet=')), None)
if not SHEET and '--sheet' in sys.argv:
    i = sys.argv.index('--sheet')
    SHEET = sys.argv[i + 1] if i + 1 < len(sys.argv) else None

# person, name on the night, the name written over it, and where
AMENDED = [
 (7519, "Ivy", "Burton", "Luke", "Penrhyn House, Clumber Road East",
  "the Register is amended to LUKE in a later hand, struck through Burton"),
]

def from_sheet(path):
    """Read the two filled-in columns off the worksheet.

    Only rows where both NAME ON THE NIGHT and AMENDED TO have been written are
    taken; everything left blank is simply a row nobody has looked at yet.
    """
    import openpyxl
    wb = openpyxl.load_workbook(os.path.expanduser(path), data_only=True)
    out, seen = [], set()
    for ws in wb.worksheets:
        head = [str(c.value or '').strip().lower() for c in ws[1]]
        try:
            col = {k: head.index(k) for k in
                   ('person', 'name as recorded', 'house', 'name on the night', 'amended to')}
        except ValueError:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            was = str(row[col['name on the night']] or '').strip()
            later = str(row[col['amended to']] or '').strip()
            pid = row[col['person']]
            if not (was and later and pid) or pid in seen:
                continue
            seen.add(pid)
            recorded = str(row[col['name as recorded']] or '').strip()
            first = was.rsplit(' ', 1)[0] if ' ' in was else recorded.rsplit(' ', 1)[0]
            surname = was.rsplit(' ', 1)[-1]
            out.append((int(pid), first, surname, later.rsplit(' ', 1)[-1],
                        str(row[col['house']] or '').strip(),
                        'the Register is amended in a later hand, struck through ' + surname))
    return out


def main():
    c = psycopg2.connect(os.environ.get('DATABASE_PUBLIC_URL') or os.environ['DATABASE_URL'])
    cur = c.cursor()
    work = from_sheet(SHEET) if SHEET else AMENDED
    if SHEET:
        print(f"  {len(work)} rows filled in on {SHEET}\n")
    for pid, fn, was, later, where, note in work:
        cur.execute("SELECT first_name, last_name, maiden_name FROM people WHERE id=%s", (pid,))
        row = cur.fetchone()
        if not row: print(f"  #{pid}: no such person"); continue
        cfn, cln, cmaiden = row
        if cln == was:
            print(f"  #{pid} {cfn} {cln}: already as the page has it"); continue
        print(f"  #{pid} {cfn} {cln} (née {cmaiden})  ->  {fn} {was}, with {later} kept as an alias")
        print(f"      {where} - {note}")
        if APPLY:
            cur.execute("""INSERT INTO person_alias (person_id, first_name, last_name, born_year, made_by)
                           SELECT %s,%s,%s,(SELECT born_year FROM people WHERE id=%s),'1939 amendment'
                           FROM (SELECT 1) t WHERE NOT EXISTS
                           (SELECT 1 FROM person_alias a WHERE a.person_id=%s AND a.last_name=%s)""",
                        (pid, fn, later, pid, pid, later))
            cur.execute("UPDATE people SET last_name=%s, maiden_name=NULL WHERE id=%s", (was, pid))
            cur.execute("""UPDATE census_entries SET source = source || %s
                            WHERE person_id=%s AND census_year=1939 AND source NOT LIKE %s""",
                        (f"; entered as {fn} {was}, which is her name on the night; {note}, and the online "
                         f"indexes show the later name", pid, '%amended to ' + later.upper() + '%'))
    print("\n  Nothing written. Add --apply." if not APPLY else "\n  done")
    if APPLY: c.commit(); print("committed")

if __name__ == '__main__':
    main()
